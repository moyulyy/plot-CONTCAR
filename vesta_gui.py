# -*- coding: utf-8 -*-
"""
CONTCAR → PPT 批量出图 GUI（iOS 风格 · PySide6）
================================================

一个本地「iOS 弹窗」风格的桌面工具：批量读取 ``CONTCAR`` / ``POSCAR``，
调用 VESTA 分别导出 **俯视图** 与 **侧视图** 截图，整理进 PPT，
并追加一页 **原子 ball + label 图例**。

所有 VESTA 相关的操作条件（坐标轴 / 晶胞边界 / 化学键 / 显示边界 /
缩放平移 / 原子半径颜色 / 视角矩阵 / 图片质量 / 隐藏窗口 / 超时等）
都做成了 **滑动开关 / 胶囊按钮 / 输入框**，无需改代码。

依赖：``PySide6``（界面）、``python-pptx``（生成 PPT）、
``Pillow`` / ``numpy``（可选，生成球体图例）。
核心出图逻辑复用 ``vesta_tools.py`` / ``vesta_modify.py``。

运行::

    D:\\miniconda3\\envs\\chem_env\\python.exe vesta_gui.py
"""

from __future__ import annotations

import html as _html
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
from math import cos, sin, radians
from pathlib import Path

# ---------------------------------------------------------------------------
# 可选第三方依赖
# ---------------------------------------------------------------------------
try:
    from PIL import Image
    import numpy as np
    _HAS_PIL = True
except Exception:  # pragma: no cover
    Image = None
    np = None
    _HAS_PIL = False

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    _HAS_PPTX = True
except Exception:
    Presentation = None
    Inches = None
    Pt = None
    RGBColor = None
    PP_ALIGN = None
    _HAS_PPTX = False

from PySide6.QtCore import (Qt, QObject, QThread, Signal, Slot, Property,
                            QPropertyAnimation, QEasingCurve, QPointF, QRectF, QTimer)
from PySide6.QtGui import QPainter, QColor, QFont, QPen, QIcon
from PySide6.QtWidgets import (
    QApplication, QWidget, QFrame, QLabel, QPushButton, QLineEdit, QComboBox,
    QSpinBox, QDoubleSpinBox, QScrollArea, QVBoxLayout, QHBoxLayout,
    QProgressBar, QPlainTextEdit, QFileDialog, QMessageBox, QColorDialog,
    QListWidget, QListWidgetItem, QTableWidget, QTableWidgetItem, QHeaderView,
    QAbstractItemView, QAbstractButton, QGraphicsDropShadowEffect, QSizePolicy,
    QRadioButton, QButtonGroup,
)

# 复用核心工具
from vesta_tools import Vesta, DEFAULT_VESTA_EXE, DEFAULT_SCALE, DEFAULT_TIMEOUT
from vesta_modify import ELEMENTS_RADIUS

# ---------------------------------------------------------------------------
# iOS 配色
# ---------------------------------------------------------------------------
BG = "#EEF0F5"
CARD = "#FFFFFF"
ACCENT = "#0A84FF"
GREEN = "#34C759"
RED = "#FF3B30"
GRAY = "#6E6E73"
LABEL = "#1C1C1E"
SECONDARY = "#5A5A60"
BORDER = "#E3E5EA"

# 程序图标（窗口 / 任务栏）
ICON_PATH = Path(__file__).resolve().parent / "assets" / "app.ico"

_ELEMENT_ORDER = ["H", "He", "Li", "Be", "B", "C", "N", "O", "F", "Ne",
                  "Na", "Mg", "Al", "Si", "P", "S", "Cl", "Ar", "K", "Ca",
                  "Sc", "Ti", "V", "Cr", "Mn", "Fe", "Co", "Ni", "Cu", "Zn",
                  "Ga", "Ge", "As", "Se", "Br", "Kr", "Rb", "Sr", "Y", "Zr",
                  "Nb", "Mo", "Tc", "Ru", "Rh", "Pd", "Ag", "Cd", "In", "Sn",
                  "Sb", "Te", "I", "Xe", "Cs", "Ba", "La", "Ce", "Pr", "Nd",
                  "Pm", "Sm", "Eu", "Gd", "Tb", "Dy", "Ho", "Er", "Tm", "Yb",
                  "Lu", "Hf", "Ta", "W", "Re", "Os", "Ir", "Pt", "Au", "Hg",
                  "Tl", "Pb", "Bi", "Po", "At", "Rn"]

# 常见元素 Jmol 配色（添加元素时的默认颜色）
JMOL_COLORS = {
    "H": "#FFFFFF", "C": "#909090", "N": "#3050F8", "O": "#FF0D0D",
    "F": "#90E050", "S": "#FFFF30", "Cl": "#1FF01F", "Br": "#A62929",
    "I": "#940094", "Fe": "#E06633", "Co": "#F090A0", "Ni": "#50D050",
    "Cu": "#C88033", "Zn": "#7D80B0", "Ti": "#BFC2C7", "V": "#A6A6AB",
    "Cr": "#8A99C7", "Mn": "#9C7AC4", "Mg": "#8AFF00", "Al": "#BFA6A6",
    "Si": "#F0C8A0", "P": "#FF8000", "K": "#8F40D4", "Ca": "#3DFF00",
    "Na": "#AB5CF2", "Li": "#CC80FF", "Mo": "#54B5B5", "W": "#2194D6",
    "Pt": "#D0D0E0", "Au": "#FFD123", "Ag": "#C0C0C0", "Pb": "#575961",
    "Zr": "#94E0E0", "Nb": "#73C2C9", "Ru": "#248F8F", "Rh": "#0A7D8C",
    "Pd": "#006985", "Cd": "#FFD98F", "Sn": "#668080", "Sb": "#9E63B5",
    "Te": "#D47A00", "Ba": "#00C900", "La": "#70D4FF", "Ce": "#FFFFC7",
}

# ---------------------------------------------------------------------------
# 工具函数
# ---------------------------------------------------------------------------
def parse_color(text: str):
    """把 '#RRGGBB' 或 'R,G,B' 解析成 'R G B'（供 modify 使用），失败返回 None。"""
    s = (text or "").strip()
    if not s:
        return None
    if s.startswith("#"):
        h = s.lstrip("#")
        if len(h) == 3:
            h = "".join(c * 2 for c in h)
        try:
            return f"{int(h[0:2], 16)} {int(h[2:4], 16)} {int(h[4:6], 16)}"
        except Exception:
            return None
    parts = [p for p in re.split(r"[,\s]+", s) if p]
    if len(parts) == 3:
        try:
            return " ".join(str(max(0, min(255, int(float(p))))) for p in parts)
        except Exception:
            return None
    return None


def parse_atomt(vesta_file: Path):
    """从 .vesta 的 ATOMT 段解析元素 -> (颜色, 半径)。返回有序列表。"""
    result = []
    try:
        text = vesta_file.read_text(encoding="utf-8")
    except Exception:
        return result
    lines = text.splitlines()
    in_atomt = False
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        first = stripped.split()[0]
        if first == "ATOMT":
            in_atomt = True
            continue
        if in_atomt and stripped[0].isalpha():
            break
        if in_atomt:
            parts = stripped.split()
            if len(parts) >= 6 and parts[0].isdigit() and parts[1] and parts[1][0].isalpha():
                try:
                    radius = float(parts[2])
                    rgb = tuple(int(round(float(x))) for x in parts[3:6])
                except Exception:
                    continue
                result.append({"element": parts[1], "radius": radius, "rgb": rgb})
    return result


def parse_contcar_elements(path) -> list:
    """从 CONTCAR/POSCAR 第 6 行的元素符号行提取去重后的元素（VASP5）。

    VASP4 格式（第 6 行为数字）无法得知元素名，返回空列表。
    """
    try:
        with open(path, "r", errors="ignore") as f:
            lines = []
            for _ in range(8):
                try:
                    lines.append(next(f))
                except StopIteration:
                    break
    except Exception:
        return []
    if len(lines) < 6:
        return []
    elems = []
    for tok in lines[5].split():
        m = re.match(r"^([A-Z][a-z]?)", tok)
        if not m:
            return []          # 数字 => VASP4，无元素行
        el = m.group(1)
        if el not in elems:
            elems.append(el)
    return elems


def make_ball_png(rgb, path: Path, size: int = 256):
    """用 PIL 生成带高光/明暗的 3D 球体 PNG（透明背景）。"""
    if Image is None or np is None:
        return None
    r, g, b = [int(c) for c in rgb]
    ss = 2  # 超采样抗锯齿
    S = size * ss
    yy, xx = np.mgrid[0:S, 0:S]
    cx = cy = (S - 1) / 2
    R = (S - 2) / 2
    dx = (xx - cx) / R
    dy = (yy - cy) / R
    d2 = dx * dx + dy * dy
    inside = d2 <= 1.0
    z = np.sqrt(np.maximum(0.0, 1.0 - d2))

    L = np.array([-0.45, -0.45, 0.77])
    L = L / np.linalg.norm(L)
    diff = np.clip(dx * L[0] + dy * L[1] + z * L[2], 0.0, 1.0)

    Sd = np.array([-0.35, -0.35, 0.87])
    Sd = Sd / np.linalg.norm(Sd)
    hx, hy, hz = dx + Sd[0], dy + Sd[1], z + Sd[2]
    hn = np.sqrt(hx * hx + hy * hy + hz * hz) + 1e-9
    spec = np.clip(hx / hn * Sd[0] + hy / hn * Sd[1] + hz / hn * Sd[2], 0, 1) ** 28

    col = np.zeros((S, S, 4), dtype=np.float64)
    col[..., 0] = np.clip(r * (0.30 + 0.70 * diff) + 255 * spec, 0, 255)
    col[..., 1] = np.clip(g * (0.30 + 0.70 * diff) + 255 * spec, 0, 255)
    col[..., 2] = np.clip(b * (0.30 + 0.70 * diff) + 255 * spec, 0, 255)

    aa = np.clip((1.0 - np.sqrt(np.minimum(d2, 1.0))) * R * 1.5, 0.0, 1.0)
    col[..., 3] = inside * aa * 255

    img = Image.fromarray(col.astype(np.uint8), "RGBA").resize((size, size), Image.LANCZOS)
    img.save(path)
    return path


# ---------------------------------------------------------------------------
# 视角矩阵
# ---------------------------------------------------------------------------
TOP = [[1, 0, 0], [0, 1, 0], [0, 0, 1]]          # 俯视图（沿 c 轴看）
SIDE_B = [[1, 0, 0], [0, 0, 1], [0, -1, 0]]      # 侧视图：沿 b 轴看（a 横，c 竖）
SIDE_A = [[0, 1, 0], [0, 0, 1], [1, 0, 0]]       # 侧视图：沿 a 轴看（b 横，c 竖）


def _matmul(A, B):
    return [[sum(A[i][k] * B[k][j] for k in range(3)) for j in range(3)] for i in range(3)]


def _rot_x(a):
    c, s = cos(radians(a)), sin(radians(a))
    return [[1, 0, 0], [0, c, -s], [0, s, c]]


def _rot_y(a):
    c, s = cos(radians(a)), sin(radians(a))
    return [[c, 0, s], [0, 1, 0], [-s, 0, c]]


def _rot_z(a):
    c, s = cos(radians(a)), sin(radians(a))
    return [[c, -s, 0], [s, c, 0], [0, 0, 1]]


def view_matrix(base, rx=0.0, ry=0.0, rz=0.0):
    """基础视角叠加屏幕轴额外旋转，返回 9 个数的 SCENE 矩阵。"""
    M = _matmul(_rot_z(rz), _matmul(_rot_y(ry), _matmul(_rot_x(rx), base)))
    return [M[i][j] for i in range(3) for j in range(3)]


def _model_bbox(im, thresh: int = 250):
    """返回图片中模型（非白区域）的外接框 (x0,y0,x1,y1)；无内容则 None。"""
    gray = im.convert("L")
    mask = gray.point(lambda v: 0 if v >= thresh else 255)
    return mask.getbbox()


def model_fraction(path, thresh: int = 250):
    """模型在图片中的最大占比（0~1），用于判断是否需要缩放/裁剪。"""
    if Image is None:
        return None
    try:
        im = Image.open(path).convert("RGB")
        bbox = _model_bbox(im, thresh)
        if not bbox:
            return None
        x0, y0, x1, y1 = bbox
        W, H = im.size
        return max((x1 - x0) / W, (y1 - y0) / H)
    except Exception:
        return None


def fit_model_in_image(path, target: float = 0.8, thresh: int = 250) -> None:
    """裁剪到模型外接框，再居中留白，使模型恰好占据整图的 target（默认 80%）。

    参考 ``mk-ppt`` 的包围盒裁剪思路：先取非白区域外接框（裁掉四周空白），
    再补一圈白边，让模型占图片的比例可控（既不顶到边、也不会太小）。
    """
    if Image is None:
        return
    try:
        im = Image.open(path).convert("RGB")
        bbox = _model_bbox(im, thresh)
        if not bbox:
            return
        x0, y0, x1, y1 = bbox
        mw, mh = x1 - x0, y1 - y0
        if mw < 4 or mh < 4:
            return
        target = min(0.98, max(0.2, float(target)))
        cw = max(mw + 1, int(round(mw / target)))
        ch = max(mh + 1, int(round(mh / target)))
        model = im.crop((x0, y0, x1, y1))
        canvas = Image.new("RGB", (cw, ch), (255, 255, 255))
        canvas.paste(model, ((cw - mw) // 2, (ch - mh) // 2))
        canvas.save(path)
    except Exception:
        pass


# ---------------------------------------------------------------------------
# PPT 生成
# ---------------------------------------------------------------------------
def build_pptx(pptx_path, slides, legend, s):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    SW, SH = 13.333, 7.5

    # 图片布局：one = 每页一张（图片约占页面 80%）；two = 俯视+侧视同页并排
    one_per_page = (s.get("slide_layout", "one") == "one")
    max_frac = float(s.get("img_frac", 0.8) or 0.8)
    max_frac = min(0.95, max(0.4, max_frac))

    def white_bg(slide):
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    def add_title(slide, text, top=0.28):
        box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(11.9), Inches(0.6))
        tf = box.text_frame
        tf.text = text
        p = tf.paragraphs[0]
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x1C, 0x1C, 0x1E)

    def add_caption(slide, text, cx, top, width, size=14):
        box = slide.shapes.add_textbox(Inches(cx - width / 2), Inches(top),
                                       Inches(width), Inches(0.4))
        tf = box.text_frame
        tf.text = text
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0x8E, 0x8E, 0x93)

    def img_size(path, bw, bh):
        """等比缩放到 (bw,bh) 框内（不超出），返回英寸 (w,h)。"""
        iw, ih = 4, 3
        if Image is not None:
            try:
                im = Image.open(path)
                iw, ih = im.size
            except Exception:
                pass
        if iw <= 0 or ih <= 0:
            iw, ih = 4, 3
        r = min(bw / iw, bh / ih)
        return iw * r, ih * r

    def img_dims(path):
        iw, ih = 4, 3
        if Image is not None:
            try:
                im = Image.open(path)
                iw, ih = im.size
            except Exception:
                pass
        if iw <= 0 or ih <= 0:
            iw, ih = 4, 3
        return iw, ih

    def put_pic(slide, path, left, top, w, h):
        # 不加边框，图片干净地放在页面上
        slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                                 Inches(w), Inches(h))

    # 标题页
    s0 = prs.slides.add_slide(blank)
    white_bg(s0)
    bar = s0.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(2.2), Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x00, 0x7A, 0xFF)
    bar.line.fill.background()
    t = s0.shapes.add_textbox(Inches(0.7), Inches(2.9), Inches(11.9), Inches(1.4))
    t.text_frame.text = "CONTCAR 结构图集"
    p = t.text_frame.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1C, 0x1C, 0x1E)
    st = s0.shapes.add_textbox(Inches(0.7), Inches(4.4), Inches(11.9), Inches(0.6))
    st.text_frame.text = (f"俯视图 / 侧视图 · 共 {len(slides)} 个结构 · "
                          f"{time.strftime('%Y-%m-%d %H:%M')}")
    st.text_frame.paragraphs[0].font.size = Pt(16)
    st.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x8E, 0x8E, 0x93)

    if s["legend"] and legend and s["legend_pos"] == "开头（标题后）":
        add_legend_slide(prs, blank, legend, s)

    # 内容区：约占页面的 max_frac（上下 / 左右各留白），图片只在其中等比缩放，绝不越界
    region_w = SW * max_frac
    region_x = (SW - region_w) / 2
    region_top = 0.90
    cap_space = 0.44
    region_h = min(SH * max_frac, SH - 0.15 - region_top - cap_space)
    region_h = max(1.0, region_h)
    img_bound_h = region_h
    cap_top = region_top + region_h + 0.04
    gap = 0.35

    def render_views(sl, views):
        n = len(views)
        if n == 1:
            path, label = views[0]
            w, h = img_size(path, region_w, region_h)
            left = region_x + (region_w - w) / 2
            top = region_top + (region_h - h) / 2
            put_pic(sl, path, left, top, w, h)
            add_caption(sl, label, region_x + region_w / 2, cap_top, region_w)
            return
        # 多图同页：统一高度（等高对齐），并约束总宽与单图宽度，避免过高 / 过宽
        aspects = [img_dims(p)[0] / img_dims(p)[1] for p, _ in views]
        total_gap = gap * (n - 1)
        avail_w = region_w - total_gap
        max_w_each = avail_w * 0.72
        h = region_h
        if sum(aspects) > 0:
            h = min(h, avail_w / sum(aspects))
        for a in aspects:
            if a > 0:
                h = min(h, max_w_each / a)
        h = max(0.4, h)
        widths = [a * h for a in aspects]
        total_w = sum(widths) + total_gap
        x = region_x + (region_w - total_w) / 2
        top = region_top + (region_h - h) / 2
        for (path, label), w in zip(views, widths):
            put_pic(sl, path, x, top, w, h)
            add_caption(sl, label, x + w / 2, cap_top, w)
            x += w + gap

    for slide in slides:
        views = []
        if slide["top"]:
            views.append((slide["top"], "俯视图 Top View"))
        if slide["side"]:
            views.append((slide["side"], "侧视图 Side View"))
        if not views:
            continue
        if one_per_page:
            for path, label in views:
                sl = prs.slides.add_slide(blank)
                white_bg(sl)
                add_title(sl, slide["name"])
                render_views(sl, [(path, label)])
        else:
            sl = prs.slides.add_slide(blank)
            white_bg(sl)
            add_title(sl, slide["name"])
            render_views(sl, views)

    if s["legend"] and legend and s["legend_pos"] == "末尾":
        add_legend_slide(prs, blank, legend, s)

    prs.save(pptx_path)


def add_legend_slide(prs, blank, legend, s):
    sl = prs.slides.add_slide(blank)
    bg = sl.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    t = sl.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11.9), Inches(0.7))
    t.text_frame.text = "原子图例 · Atom Legend"
    p = t.text_frame.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1C, 0x1C, 0x1E)

    ball_in = max(0.3, s["ball_cm"] / 2.54)
    cols = max(1, s["legend_cols"])
    cell_w = 11.9 / cols
    start_y = 1.5
    cell_h = ball_in + 0.75

    tmp = Path(tempfile.mkdtemp(prefix="legend_balls_"))
    try:
        for i, item in enumerate(legend):
            row, col = divmod(i, cols)
            cx = 0.7 + col * cell_w + cell_w / 2
            cy = start_y + row * cell_h

            ball_path = tmp / f"{i}_{item['element']}.png"
            make_ball_png(item["rgb"], ball_path, size=256)
            if ball_path.exists():
                sl.shapes.add_picture(str(ball_path), Inches(cx - ball_in / 2),
                                      Inches(cy), Inches(ball_in), Inches(ball_in))
            else:
                from pptx.enum.shapes import MSO_SHAPE
                shp = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - ball_in / 2),
                                          Inches(cy), Inches(ball_in), Inches(ball_in))
                shp.fill.solid()
                shp.fill.fore_color.rgb = RGBColor(*item["rgb"])
                shp.line.fill.background()

            label = item["element"]
            if s["legend_radius"]:
                label += f"\nr={item['radius']:.2f} Å"
            lb = sl.shapes.add_textbox(Inches(cx - cell_w / 2),
                                       Inches(cy + ball_in + 0.05),
                                       Inches(cell_w), Inches(0.8))
            tf = lb.text_frame
            tf.text = label
            tf.word_wrap = True
            for para in tf.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                para.font.size = Pt(14)
                para.font.color.rgb = RGBColor(0x1C, 0x1C, 0x1E)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return sl


# ---------------------------------------------------------------------------
# 批量处理（无界面依赖）
# ---------------------------------------------------------------------------
def run_batch(s, log, progress, stop_event, pause_event=None):
    """执行批量出图 + PPT 生成。log(level,msg) / progress(0..1) 为回调。"""
    out_dir = Path(s["out_dir"]).expanduser()
    img_dir = out_dir / "images"
    vesta_dir = out_dir / "vesta"
    out_dir.mkdir(parents=True, exist_ok=True)
    img_dir.mkdir(parents=True, exist_ok=True)
    if s["keep_vesta"]:
        vesta_dir.mkdir(parents=True, exist_ok=True)

    pptx_path = Path(s["ppt_path"]).expanduser()
    if pptx_path.suffix.lower() != ".pptx":
        pptx_path = pptx_path.with_suffix(".pptx")
    pptx_path.parent.mkdir(parents=True, exist_ok=True)

    v = Vesta(exe=s["vesta_exe"], timeout=s["timeout"], verbose=False,
              show_window=not s["hide_window"])

    tmpdir = Path(tempfile.mkdtemp(prefix="contcar_ppt_"))
    try:
        log("info", f"VESTA 路径：{s['vesta_exe']}")
        log("info", f"输出目录：{out_dir}")
        log("info", f"PPT 路径：{pptx_path}")
        log("info", f"共 {len(s['files'])} 个结构文件。")

        top_mat = view_matrix(TOP, s["rot_x"], s["rot_y"], s["rot_z"])
        side_mat = view_matrix(SIDE_A if s["side_dir"] == "沿 a 轴看" else SIDE_B,
                               s["rot_x"], s["rot_y"], s["rot_z"])

        base_modify = {
            "comps": s["comps"], "ucolp": s["ucolp"], "sbond": s["sbond"],
            "sbond_padding": s["sbond_padding"],
        }
        if s["boundary"]:
            base_modify["boundary"] = s["boundary"]
        if s["atom_params"]:
            base_modify["atom_params"] = s["atom_params"]

        slides = []
        legend = []
        n_files = len(s["files"])
        views_per = int(bool(s["do_top"])) + int(bool(s["do_side"]))
        total = n_files * max(1, views_per) + 1
        done = 0
        ext = s["img_fmt"].lstrip(".")

        def wait_pause():
            """暂停时在此等待（可被停止打断）。"""
            while (pause_event is not None and pause_event.is_set()
                   and not stop_event.is_set()):
                time.sleep(0.2)

        for idx, item in enumerate(s["files"]):
            if stop_event.is_set():
                return False, "已停止（未生成完整 PPT）。"
            wait_pause()
            src = Path(item["path"])
            rel_sub = item.get("rel", "")            # 相对子目录（posix，可为空）
            stem = item.get("stem") or src.stem or "structure"
            label = item.get("label") or stem
            if not src.exists():
                log("warn", f"[跳过] 文件不存在：{src}")
                done += views_per
                progress(done / total)
                continue

            log("info", f"[{idx+1}/{n_files}] 处理 {label}")
            try:
                base_vesta = tmpdir / f"{idx}.vesta"
                if src.suffix.lower() == ".vesta":
                    shutil.copy(src, base_vesta)          # 已是 .vesta，直接使用
                else:
                    v.contcar_to_vesta(src, base_vesta)
                if base_modify:
                    v.modify_vesta(base_vesta, **base_modify)

                for it in parse_atomt(base_vesta):
                    if it["element"] not in [x["element"] for x in legend]:
                        legend.append(it)

                # 保持相对目录结构
                img_sub = img_dir / rel_sub if rel_sub else img_dir
                img_sub.mkdir(parents=True, exist_ok=True)
                vesta_sub = vesta_dir / rel_sub if rel_sub else vesta_dir
                if s["keep_vesta"]:
                    vesta_sub.mkdir(parents=True, exist_ok=True)

                slide = {"name": label, "top": None, "side": None}
                target = float(s.get("model_frac", 0.8) or 0.8)
                fit_mode = s.get("model_fit", "crop")   # crop / zoom / both

                def process_view(vf, img, mat):
                    """导出并让模型在图片中约占 target（VESTA 缩放 / 裁剪 或两者）。"""
                    base_frac = float(s["scale_frac"]) or 1.0
                    v.modify_vesta(vf, version=mat, x_move_frac=s["x_move"],
                                   y_move_frac=s["y_move"], scale_frac=base_frac)
                    v.export_image(vf, img, scale=s["scale"])
                    if fit_mode in ("zoom", "both"):
                        f = model_fraction(img)
                        if f and f > 1e-6:
                            new_frac = max(0.05, min(50.0, base_frac * target / f))
                            if abs(new_frac - base_frac) > 0.02 * base_frac:
                                v.modify_vesta(vf, version=mat, x_move_frac=s["x_move"],
                                               y_move_frac=s["y_move"], scale_frac=new_frac)
                                v.export_image(vf, img, scale=s["scale"])
                    if fit_mode in ("crop", "both"):
                        fit_model_in_image(img, target)

                if s["do_top"]:
                    vf = tmpdir / f"{idx}_top.vesta"
                    shutil.copy(base_vesta, vf)
                    img = img_sub / f"{stem}_top.{ext}"
                    process_view(vf, img, top_mat)
                    if s["keep_vesta"]:
                        shutil.copy(vf, vesta_sub / f"{stem}_top.vesta")
                    slide["top"] = img
                    log("ok", f"    俯视图 → {img.relative_to(out_dir)}")
                    done += 1
                    progress(done / total)

                wait_pause()
                if s["do_side"] and not stop_event.is_set():
                    vf = tmpdir / f"{idx}_side.vesta"
                    shutil.copy(base_vesta, vf)
                    img = img_sub / f"{stem}_side.{ext}"
                    process_view(vf, img, side_mat)
                    if s["keep_vesta"]:
                        shutil.copy(vf, vesta_sub / f"{stem}_side.vesta")
                    slide["side"] = img
                    log("ok", f"    侧视图 → {img.relative_to(out_dir)}")
                    done += 1
                    progress(done / total)

                slides.append(slide)
            except Exception as exc:  # noqa: BLE001
                log("err", f"    [失败] {exc}")

        if stop_event.is_set():
            return False, "已停止（未生成完整 PPT）。"

        log("info", "正在生成 PPT…")
        build_pptx(pptx_path, slides, legend, s)
        done += 1
        progress(done / total)

        return True, f"完成！PPT 已保存：{pptx_path}"
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


# ---------------------------------------------------------------------------
# QSS 样式（iOS 风格）
# ---------------------------------------------------------------------------
QSS = """
* { font-family: 'Segoe UI', 'Microsoft YaHei UI', 'PingFang SC'; outline: none; }
QWidget#Root {
    background: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 #F8F9FC, stop:1 #ECEEF4);
    border-radius: 16px;
}

QLabel#Title { font-size: 21px; font-weight: 700; color: #1C1C1E; background: transparent; }
QLabel#Subtitle { font-size: 12px; color: #5A5A60; background: transparent; }
QLabel#SectionLabel { font-size: 12px; font-weight: 600; color: #6E6E73; background: transparent; }
QLabel#CardTitle { font-size: 14px; font-weight: 600; color: #1C1C1E; background: transparent; }
QLabel#CardSubtitle { font-size: 11px; color: #6E6E73; background: transparent; }
QLabel#FieldLabel { font-size: 13px; color: #1C1C1E; background: transparent; }
QLabel#HintLabel { font-size: 11px; color: #6E6E73; background: transparent; }

QRadioButton { color: #1C1C1E; font-size: 13px; background: transparent; spacing: 6px; }
QRadioButton::indicator { width: 15px; height: 15px; }

QFrame#Card { background: #FFFFFF; border: 1px solid #E3E5EA; border-radius: 14px; }

QLineEdit, QSpinBox, QDoubleSpinBox, QComboBox {
    background: #F8F8FA; border: 1px solid #D9D9DE; border-radius: 8px;
    padding: 5px 8px; font-size: 13px; color: #1C1C1E;
    selection-background-color: #0A84FF; selection-color: #FFFFFF;
}
QLineEdit:focus, QSpinBox:focus, QDoubleSpinBox:focus, QComboBox:focus {
    border: 1px solid #0A84FF; background: #FFFFFF;
}
QLineEdit:hover, QSpinBox:hover, QDoubleSpinBox:hover, QComboBox:hover { border: 1px solid #B9BCC4; }
QSpinBox::up-button, QSpinBox::down-button,
QDoubleSpinBox::up-button, QDoubleSpinBox::down-button {
    width: 0px; height: 0px; border: none; background: transparent;
}
QComboBox::drop-down { border: none; width: 24px; }
QComboBox QAbstractItemView {
    background: #FFFFFF; border: 1px solid #E3E5EA; border-radius: 10px;
    padding: 4px; selection-background-color: #EAF3FF; selection-color: #0A84FF;
}

QListWidget { background: #F8F8FA; border: 1px solid #D9D9DE; border-radius: 10px;
              padding: 4px; font-size: 12px; color: #1C1C1E; }
QListWidget::item { padding: 4px 6px; border-radius: 6px; }
QListWidget::item:selected { background: #EAF3FF; color: #0A84FF; }
QListWidget::item:hover { background: #F2F2F7; }

QTableWidget { background: #F8F8FA; border: 1px solid #D9D9DE; border-radius: 10px;
               gridline-color: #ECECEF; font-size: 12px; color: #1C1C1E; }
QTableWidget::item { padding: 2px 6px; }
QTableWidget::item:selected { background: #EAF3FF; color: #0A84FF; }
QHeaderView::section { background: #FFFFFF; color: #6E6E73; border: none;
                       border-bottom: 1px solid #E5E5EA; padding: 6px; font-weight: 600; }

QPushButton#Pill { border: none; border-radius: 17px; font-size: 13px; font-weight: 600; padding: 0px 20px; }
QPushButton#Pill[variant="accent"] { background: #0A84FF; color: #FFFFFF; }
QPushButton#Pill[variant="accent"]:hover { background: #3395FF; }
QPushButton#Pill[variant="accent"]:pressed { background: #006EDB; }
QPushButton#Pill[variant="green"] { background: #34C759; color: #FFFFFF; }
QPushButton#Pill[variant="green"]:hover { background: #5DD57B; }
QPushButton#Pill[variant="green"]:pressed { background: #28A745; }
QPushButton#Pill[variant="red"] { background: #FF3B30; color: #FFFFFF; }
QPushButton#Pill[variant="red"]:hover { background: #FF655C; }
QPushButton#Pill[variant="red"]:pressed { background: #D92B20; }
QPushButton#Pill[variant="gray"] { background: #E9E9EB; color: #1C1C1E; }
QPushButton#Pill[variant="gray"]:hover { background: #DFDFE3; }
QPushButton#Pill[variant="gray"]:pressed { background: #D0D0D4; }
QPushButton#Pill[variant="outline"] { background: #FFFFFF; color: #0A84FF; border: 1px solid #0A84FF; }
QPushButton#Pill[variant="outline"]:hover { background: #EAF3FF; }
QPushButton#Pill[variant="outline"]:pressed { background: #D6E9FF; }
QPushButton#Pill:disabled { background: #E5E5EA; color: #8E8E93; border: none; }

QPushButton#WinBtn { background: #E9E9EB; color: #3C3C43; border: none; border-radius: 14px; font-weight: 700; }
QPushButton#WinBtn:hover { background: #DFDFE3; }
QPushButton#CloseBtn { background: #FF3B30; color: #FFFFFF; border: none; border-radius: 14px; font-weight: 700; }
QPushButton#CloseBtn:hover { background: #FF655C; }
QPushButton#DeleteBtn { background: transparent; color: #FF3B30; border: none; border-radius: 14px; font-weight: 700; }
QPushButton#DeleteBtn:hover { background: #FCE8E6; }

QProgressBar { background: #E5E5EA; border: none; border-radius: 5px; height: 10px; }
QProgressBar::chunk { background: #34C759; border-radius: 5px; }

QPlainTextEdit#Log { background: #1C1C1E; color: #FFFFFF; border: none; border-radius: 10px;
                     padding: 10px; font-family: 'Consolas', 'Courier New'; font-size: 12px; }

QScrollArea { background: transparent; border: none; }
QScrollArea > QWidget > QWidget { background: transparent; }
QScrollBar:vertical { background: transparent; width: 8px; margin: 2px; }
QScrollBar::handle:vertical { background: #C7C7CC; border-radius: 4px; min-height: 30px; }
QScrollBar::handle:vertical:hover { background: #AEAEB2; }
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0px; width: 0px; }
QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical { background: transparent; }

QPlainTextEdit QScrollBar:vertical { background: #2C2C2E; width: 8px; }
QPlainTextEdit QScrollBar::handle:vertical { background: #5A5A5E; }
QPlainTextEdit QScrollBar::handle:vertical:hover { background: #6E6E73; }
QPlainTextEdit QScrollBar::add-line:vertical, QPlainTextEdit QScrollBar::sub-line:vertical { height: 0px; }
"""


# ---------------------------------------------------------------------------
# iOS 开关（自绘 + 动画）
# ---------------------------------------------------------------------------
class Toggle(QAbstractButton):
    """iOS 风格滑动开关。"""

    def __init__(self, checked=False, parent=None):
        super().__init__(parent)
        self.setCheckable(True)
        self.setCursor(Qt.PointingHandCursor)
        self.setFixedSize(50, 30)
        self._thumb = 2.0 if not checked else 24.0
        self.setChecked(checked)
        self.toggled.connect(self._on_toggled)
        self._anim = None

    def _get_thumb(self) -> float:
        return self._thumb

    def _set_thumb(self, v: float):
        self._thumb = v
        self.update()

    thumbPos = Property(float, _get_thumb, _set_thumb)

    def _on_toggled(self, checked):
        anim = QPropertyAnimation(self, b"thumbPos", self)
        anim.setDuration(160)
        anim.setStartValue(self._thumb)
        anim.setEndValue(2.0 if not checked else 24.0)
        anim.setEasingCurve(QEasingCurve.OutCubic)
        anim.start()
        self._anim = anim

    def paintEvent(self, event):
        p = QPainter(self)
        p.setRenderHint(QPainter.Antialiasing)
        w, h = self.width(), self.height()
        track = QColor(GREEN) if self.isChecked() else QColor("#D1D1D6")
        p.setPen(Qt.NoPen)
        p.setBrush(track)
        p.drawRoundedRect(QRectF(0, h / 2 - 13, w, 26), 13, 13)
        # 滑块阴影 + 滑块
        p.setBrush(QColor(0, 0, 0, 18))
        p.drawEllipse(QPointF(self._thumb + 13, h / 2 + 1), 11.5, 11.5)
        p.setBrush(QColor("#FFFFFF"))
        p.drawEllipse(QPointF(self._thumb + 13, h / 2), 11, 11)


class NoWheelComboBox(QComboBox):
    """下拉框：忽略滚轮，避免滚动页面时误改选项。"""

    def wheelEvent(self, e):
        e.ignore()


class NoWheelSpinBox(QSpinBox):
    """整数输入框：忽略滚轮。"""

    def wheelEvent(self, e):
        e.ignore()


class NoWheelDoubleSpinBox(QDoubleSpinBox):
    """浮点输入框：忽略滚轮。"""

    def wheelEvent(self, e):
        e.ignore()


class Card(QFrame):
    """圆角白色卡片（带阴影）。"""

    def __init__(self, title=None, subtitle=None, parent=None):
        super().__init__(parent)
        self.setObjectName("Card")
        shadow = QGraphicsDropShadowEffect(self)
        shadow.setBlurRadius(22)
        shadow.setOffset(0, 4)
        shadow.setColor(QColor(0, 0, 0, 28))
        self.setGraphicsEffect(shadow)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)

        if title:
            t = QLabel(title)
            t.setObjectName("CardTitle")
            t.setContentsMargins(18, 16, 18, 0)
            lay.addWidget(t)
        if subtitle:
            st = QLabel(subtitle)
            st.setObjectName("CardSubtitle")
            st.setContentsMargins(18, 2, 18, 8)
            st.setWordWrap(True)
            lay.addWidget(st)
        if title or subtitle:
            sep_wrap = QWidget()
            sep_lay = QHBoxLayout(sep_wrap)
            sep_lay.setContentsMargins(18, 0, 18, 0)
            sep = QFrame()
            sep.setFixedHeight(1)
            sep.setStyleSheet("background: #E5E5EA;")
            sep_lay.addWidget(sep)
            lay.addWidget(sep_wrap)


class TitleBar(QWidget):
    """可拖动的顶部标题栏 + 窗口控制按钮。"""

    def __init__(self, window):
        super().__init__(window)
        self._window = window
        self._drag_offset = None
        self.setFixedHeight(66)

        lay = QHBoxLayout(self)
        lay.setContentsMargins(24, 12, 14, 8)
        lay.setSpacing(6)

        col = QVBoxLayout()
        col.setSpacing(1)
        t = QLabel("CONTCAR → PPT")
        t.setObjectName("Title")
        st = QLabel("批量导出俯视图 / 侧视图，整理成 PPT，并生成原子图例页")
        st.setObjectName("Subtitle")
        col.addWidget(t)
        col.addWidget(st)
        lay.addLayout(col)
        lay.addStretch(1)

        min_btn = QPushButton("—")
        min_btn.setObjectName("WinBtn")
        min_btn.setFixedSize(28, 28)
        min_btn.setCursor(Qt.PointingHandCursor)
        min_btn.clicked.connect(window.showMinimized)

        close_btn = QPushButton("✕")
        close_btn.setObjectName("CloseBtn")
        close_btn.setFixedSize(28, 28)
        close_btn.setCursor(Qt.PointingHandCursor)
        close_btn.clicked.connect(window.close)

        lay.addWidget(min_btn)
        lay.addWidget(close_btn)

    def mousePressEvent(self, e):
        if e.button() == Qt.LeftButton:
            self._drag_offset = e.globalPosition().toPoint() - self._window.frameGeometry().topLeft()

    def mouseMoveEvent(self, e):
        if self._drag_offset is not None and e.buttons() & Qt.LeftButton:
            self._window.move(e.globalPosition().toPoint() - self._drag_offset)

    def mouseReleaseEvent(self, e):
        self._drag_offset = None


# ---------------------------------------------------------------------------
# 后台工作线程
# ---------------------------------------------------------------------------
class Worker(QObject):
    log = Signal(str, str)
    progress = Signal(float)
    done = Signal(bool, str)

    def __init__(self, settings, stop_event, pause_event=None):
        super().__init__()
        self.settings = settings
        self.stop_event = stop_event
        self.pause_event = pause_event

    @Slot()
    def run(self):
        try:
            ok, msg = run_batch(self.settings, self._log, self._progress,
                                self.stop_event, self.pause_event)
        except Exception as exc:  # noqa: BLE001
            self._log("err", f"发生错误：{exc}")
            self._log("err", traceback.format_exc())
            ok, msg = False, f"处理失败：{exc}"
        self.done.emit(ok, msg)

    def _log(self, level, msg):
        self.log.emit(level, msg)

    def _progress(self, f):
        self.progress.emit(f)


# ---------------------------------------------------------------------------
# 主窗口
# ---------------------------------------------------------------------------
class MainWindow(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowFlags(Qt.FramelessWindowHint | Qt.Window
                            | Qt.WindowStaysOnTopHint)
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.resize(968, 830)

        self.stop_event = threading.Event()
        self.pause_event = threading.Event()
        self.running = False
        self.thread = None
        self.worker = None
        self.root_path = None
        self.atom_rows: list[dict] = []

        # 外层透明留白（给窗口阴影）
        outer = QVBoxLayout(self)
        outer.setContentsMargins(26, 26, 26, 26)
        self.root = QWidget()
        self.root.setObjectName("Root")
        shadow = QGraphicsDropShadowEffect(self.root)
        shadow.setBlurRadius(28)
        shadow.setOffset(0, 6)
        shadow.setColor(QColor(0, 0, 0, 60))
        self.root.setGraphicsEffect(shadow)
        outer.addWidget(self.root)

        root_lay = QVBoxLayout(self.root)
        root_lay.setContentsMargins(0, 0, 0, 0)
        root_lay.setSpacing(0)

        self.titlebar = TitleBar(self)
        root_lay.addWidget(self.titlebar)

        # 滚动区
        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.scroll.setFrameShape(QFrame.NoFrame)
        self.scroll.viewport().setAutoFillBackground(False)
        self.body = QWidget()
        self.body.setAutoFillBackground(False)
        self.scroll.setWidget(self.body)
        self.body_lay = QVBoxLayout(self.body)
        self.body_lay.setContentsMargins(18, 4, 18, 26)
        self.body_lay.setSpacing(12)
        root_lay.addWidget(self.scroll, 1)

        self._build_sections()

    # ------------------------------------------------------------------ #
    # 控件工厂
    # ------------------------------------------------------------------ #
    def _line(self, default, width=80, name=None):
        e = QLineEdit(default)
        e.setFixedWidth(width)
        if name:
            setattr(self, name, e)
        return e

    def _dspin(self, default, lo, hi, width=84, dec=1, name=None):
        s = NoWheelDoubleSpinBox()
        s.setRange(lo, hi)
        s.setDecimals(dec)
        s.setValue(float(default))
        s.setFixedWidth(width)
        s.setKeyboardTracking(False)
        if name:
            setattr(self, name, s)
        return s

    def _spin(self, default, lo, hi, width=84, name=None):
        s = NoWheelSpinBox()
        s.setRange(lo, hi)
        s.setValue(int(default))
        s.setFixedWidth(width)
        if name:
            setattr(self, name, s)
        return s

    def _combo(self, values, width=130, name=None):
        c = NoWheelComboBox()
        c.addItems(values)
        c.setFixedWidth(width)
        if name:
            setattr(self, name, c)
        return c

    def _toggle(self, default, name=None):
        t = Toggle(checked=default)
        if name:
            setattr(self, name, t)
        return t

    def _pill(self, text, variant, height=34, name=None):
        b = QPushButton(text)
        b.setObjectName("Pill")
        b.setProperty("variant", variant)
        b.setCursor(Qt.PointingHandCursor)
        b.setFixedHeight(height)
        if name:
            setattr(self, name, b)
        return b

    def _section_label(self, text):
        lbl = QLabel(text)
        lbl.setObjectName("SectionLabel")
        self.body_lay.addWidget(lbl)

    def _card(self, title=None, subtitle=None):
        card = Card(title, subtitle)
        self.body_lay.addWidget(card)
        return card

    def _row(self, card, label, widget):
        h = QHBoxLayout()
        h.setContentsMargins(18, 6, 18, 6)
        lbl = QLabel(label)
        lbl.setObjectName("FieldLabel")
        h.addWidget(lbl)
        h.addStretch(1)
        h.addWidget(widget, 0, Qt.AlignVCenter)
        card.layout().addLayout(h)

    # ------------------------------------------------------------------ #
    # 各分区
    # ------------------------------------------------------------------ #
    def _build_sections(self):
        # ① 输入文件
        self._section_label("①  输入文件（按文件夹检索 CONTCAR / .vesta）")
        card = self._card("结构根目录",
                          "给一个根目录，递归检索其下所有结构文件；勾选要处理的项，"
                          "输出时保持相对路径结构。")
        in_row = QHBoxLayout()
        in_row.setContentsMargins(18, 10, 18, 2)
        in_row.setSpacing(8)
        self.input_root = QLineEdit()
        self.input_root.setPlaceholderText("选择或粘贴一个文件夹路径…")
        in_row.addWidget(self.input_root, 1)
        b_in = self._pill("浏览", "gray", height=30)
        b_in.clicked.connect(self._browse_input)
        in_row.addWidget(b_in)
        b_scan = self._pill("检索", "accent", height=30)
        b_scan.clicked.connect(self._scan)
        in_row.addWidget(b_scan)
        card.layout().addLayout(in_row)

        type_row = QHBoxLayout()
        type_row.setContentsMargins(18, 2, 18, 4)
        type_row.setSpacing(14)
        tlbl = QLabel("检索类型")
        tlbl.setObjectName("FieldLabel")
        type_row.addWidget(tlbl)
        self.radio_contcar = QRadioButton("CONTCAR 文件")
        self.radio_vesta = QRadioButton(".vesta 文件")
        self.radio_contcar.setChecked(True)
        self.radio_group = QButtonGroup(self)
        self.radio_group.addButton(self.radio_contcar)
        self.radio_group.addButton(self.radio_vesta)
        self.radio_contcar.toggled.connect(self._on_type_changed)
        self.radio_vesta.toggled.connect(self._on_type_changed)
        type_row.addWidget(self.radio_contcar)
        type_row.addWidget(self.radio_vesta)
        type_row.addStretch(1)
        card.layout().addLayout(type_row)

        tool_row = QHBoxLayout()
        tool_row.setContentsMargins(18, 4, 18, 2)
        tool_row.setSpacing(8)
        b_all = self._pill("全选", "outline", height=28)
        b_all.clicked.connect(lambda: self._check_all(True))
        b_none = self._pill("全不选", "outline", height=28)
        b_none.clicked.connect(lambda: self._check_all(False))
        b_ref = self._pill("按选中刷新元素", "outline", height=28)
        b_ref.clicked.connect(self._refresh_elements)
        b_up = self._pill("↑ 上移", "outline", height=28)
        b_up.clicked.connect(lambda: self._move_row(-1))
        b_dn = self._pill("↓ 下移", "outline", height=28)
        b_dn.clicked.connect(lambda: self._move_row(1))
        b_top = self._pill("⤒ 置顶", "outline", height=28)
        b_top.clicked.connect(lambda: self._move_to(False))
        b_bot = self._pill("⤓ 置底", "outline", height=28)
        b_bot.clicked.connect(lambda: self._move_to(True))
        tool_row.addWidget(b_all)
        tool_row.addWidget(b_none)
        tool_row.addWidget(b_ref)
        tool_row.addWidget(b_up)
        tool_row.addWidget(b_dn)
        tool_row.addWidget(b_top)
        tool_row.addWidget(b_bot)
        tool_row.addStretch(1)
        self.scan_info = QLabel("尚未检索")
        self.scan_info.setObjectName("HintLabel")
        tool_row.addWidget(self.scan_info)
        card.layout().addLayout(tool_row)

        self.file_table = QTableWidget(0, 2)
        self.file_table.setHorizontalHeaderLabels(["结构（勾选）", "PPT 页面标题（可编辑）"])
        self.file_table.verticalHeader().setVisible(False)
        self.file_table.setFixedHeight(330)
        self.file_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.file_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.file_table.setSortingEnabled(True)
        self.file_table.setEditTriggers(QAbstractItemView.DoubleClicked
                                        | QAbstractItemView.SelectedClicked
                                        | QAbstractItemView.EditKeyPressed)
        hh = self.file_table.horizontalHeader()
        hh.setSectionResizeMode(0, QHeaderView.Stretch)
        hh.setSectionResizeMode(1, QHeaderView.Stretch)
        hh.setFixedHeight(34)
        hh.setSortIndicatorShown(True)
        hh.sectionClicked.connect(self._on_header_clicked)
        vh = self.file_table.verticalHeader()
        vh.setDefaultSectionSize(36)
        vh.setMinimumSectionSize(36)
        list_wrap = QWidget()
        list_lay = QHBoxLayout(list_wrap)
        list_lay.setContentsMargins(18, 4, 18, 14)
        list_lay.addWidget(self.file_table)
        card.layout().addWidget(list_wrap)

        # ② 输出设置
        self._section_label("②  输出设置")
        card = self._card("输出目录与 PPT",
                          "图片 → images/ 子目录，.vesta → vesta/ 子目录，均保持相对路径结构。")
        out_row = QHBoxLayout()
        out_row.setContentsMargins(18, 12, 18, 4)
        lbl = QLabel("输出目录")
        lbl.setObjectName("FieldLabel")
        out_row.addWidget(lbl)
        self.out_dir = QLineEdit(str(Path.cwd() / "ppt_output"))
        out_row.addWidget(self.out_dir, 1)
        browse = self._pill("浏览", "gray", height=30)
        browse.clicked.connect(self._browse_outdir)
        out_row.addWidget(browse)
        card.layout().addLayout(out_row)

        ppt_row = QHBoxLayout()
        ppt_row.setContentsMargins(18, 4, 18, 4)
        lbl = QLabel("PPT 保存路径")
        lbl.setObjectName("FieldLabel")
        ppt_row.addWidget(lbl)
        self.ppt_path = QLineEdit(str(Path.cwd() / "ppt_output" / "structures.pptx"))
        ppt_row.addWidget(self.ppt_path, 1)
        pb = self._pill("浏览", "gray", height=30)
        pb.clicked.connect(self._browse_ppt)
        ppt_row.addWidget(pb)
        card.layout().addLayout(ppt_row)

        self._row(card, "图片格式", self._combo(["png", "jpg", "bmp", "tif"], 110, "img_fmt_var"))
        self._row(card, "图片质量 scale", self._spin(DEFAULT_SCALE, 1, 12, 84, "scale_var"))
        self._row(card, "单个任务超时 (秒)", self._spin(DEFAULT_TIMEOUT, 30, 3600, 84, "timeout_var"))
        self._row(card, "隐藏 VESTA 窗口", self._toggle(True, "hide_window_var"))
        self._row(card, "输出 .vesta 文件", self._toggle(True, "keep_vesta_var"))

        vesta_row = QHBoxLayout()
        vesta_row.setContentsMargins(18, 4, 18, 14)
        lbl = QLabel("VESTA.exe 路径")
        lbl.setObjectName("FieldLabel")
        vesta_row.addWidget(lbl)
        self.vesta_exe_var = QLineEdit(DEFAULT_VESTA_EXE)
        vesta_row.addWidget(self.vesta_exe_var, 1)
        vbrowse = self._pill("浏览", "gray", height=30)
        vbrowse.clicked.connect(self._browse_vesta)
        vesta_row.addWidget(vbrowse)
        card.layout().addLayout(vesta_row)

        # ③ 视图设置
        self._section_label("③  视图设置")
        card = self._card("俯视图 / 侧视图")
        self._row(card, "生成俯视图（沿 c 轴）", self._toggle(True, "do_top_var"))
        self._row(card, "生成侧视图（c 轴竖直）", self._toggle(True, "do_side_var"))
        self._row(card, "侧视方向", self._combo(["沿 b 轴看", "沿 a 轴看"], 130, "side_dir_var"))
        self._row(card, "额外旋转 X (°)", self._dspin(0, -180, 180, 84, 0, "rot_x_var"))
        self._row(card, "额外旋转 Y (°)", self._dspin(0, -180, 180, 84, 0, "rot_y_var"))
        self._row(card, "额外旋转 Z (°)", self._dspin(0, -180, 180, 84, 0, "rot_z_var"))
        self._row(card, "模型占图片比例", self._dspin(0.8, 0.3, 0.95, 84, 2, "model_frac_var"))
        self._row(card, "模型放大方式", self._combo(
            ["裁剪留白（推荐）", "VESTA 缩放", "裁剪 + VESTA 缩放"], 180, "model_fit_var"))
        self._row(card, "图片排版", self._combo(
            ["同页并排（俯视+侧视）", "每页一张（约占页面 80%）"], 230, "slide_layout_var"))
        self._row(card, "图片占页面比例", self._dspin(0.8, 0.4, 0.95, 84, 2, "img_frac_var"))

        # ④ VESTA 显示
        self._section_label("④  VESTA 显示选项")
        card = self._card("结构显示（对应 .vesta 各段）")
        self._row(card, "显示晶胞坐标轴 COMPS", self._toggle(False, "comps_var"))
        self._row(card, "显示晶胞边界 UCOLP", self._toggle(True, "ucolp_var"))
        self._row(card, "生成化学键 SBOND", self._toggle(True, "sbond_var"))
        self._row(card, "成键容差 padding (Å)", self._line("0.5", 84, "sbond_padding_var"))
        self._row(card, "场景缩放 scale_frac", self._line("1.0", 84, "scale_frac_var"))
        self._row(card, "水平平移 x_move", self._line("0.0", 84, "x_move_var"))
        self._row(card, "垂直平移 y_move", self._line("0.0", 84, "y_move_var"))

        bnd_wrap = QWidget()
        bnd_lay = QVBoxLayout(bnd_wrap)
        bnd_lay.setContentsMargins(18, 4, 18, 12)
        bnd_lbl = QLabel("显示边界 BOUND（分数坐标区间；a←→x、b←→y、c←→z，留空 = 不改）")
        bnd_lbl.setObjectName("FieldLabel")
        bnd_lay.addWidget(bnd_lbl)
        bnd_grid = QHBoxLayout()
        bnd_grid.setSpacing(8)
        self.boundary_vars = []
        for axis, dflt in (("a", ("-0.05", "1.05")),
                           ("b", ("-0.05", "1.05")),
                           ("c", ("-0.05", "0.95"))):
            tag = QLabel(f"{axis} 轴")
            tag.setObjectName("HintLabel")
            tag.setFixedWidth(30)
            bnd_grid.addWidget(tag)
            for d in dflt:
                e = QLineEdit(d)
                e.setFixedWidth(58)
                e.setAlignment(Qt.AlignCenter)
                self.boundary_vars.append(e)
                bnd_grid.addWidget(e)
            if axis != "c":
                sp = QLabel("   ")
                sp.setFixedWidth(10)
                bnd_grid.addWidget(sp)
        bnd_grid.addStretch(1)
        bnd_lay.addLayout(bnd_grid)
        card.layout().addWidget(bnd_wrap)

        # ⑤ 原子样式
        self._section_label("⑤  原子半径 / 颜色（留空则用 VESTA 默认）")
        card = self._card("原子样式 ATOMT / SITET",
                          "只对列出的元素生效；未列出的元素保持 VESTA 默认样式。")
        self._row(card, "启用原子样式处理", self._toggle(True, "apply_atom_style_var"))
        head = QHBoxLayout()
        head.setContentsMargins(18, 8, 18, 0)
        head.setSpacing(6)
        for text, w in [("元素", 70), ("半径 (Å)", 80), ("颜色", 100)]:
            h = QLabel(text)
            h.setObjectName("HintLabel")
            h.setFixedWidth(w)
            head.addWidget(h)
        head.addStretch(1)
        card.layout().addLayout(head)

        self.atom_box = QWidget()
        self.atom_box_lay = QVBoxLayout(self.atom_box)
        self.atom_box_lay.setContentsMargins(18, 2, 18, 4)
        self.atom_box_lay.setSpacing(4)
        card.layout().addWidget(self.atom_box)

        add_row = QHBoxLayout()
        add_row.setContentsMargins(18, 2, 18, 14)
        self.element_combo = NoWheelComboBox()
        self.element_combo.addItems(_ELEMENT_ORDER)
        self.element_combo.setCurrentText("Fe")
        self.element_combo.setFixedWidth(90)
        add_row.addWidget(self.element_combo)
        add_btn = self._pill("＋ 添加元素", "outline")
        add_btn.clicked.connect(lambda: self._add_atom_row())
        add_row.addWidget(add_btn)
        add_row.addStretch(1)
        card.layout().addLayout(add_row)

        # ⑥ 图例页
        self._section_label("⑥  图例页（原子 ball + label）")
        card = self._card("输出原子图例 PPT 页面")
        self._row(card, "生成图例页", self._toggle(True, "legend_var"))
        self._row(card, "球体直径 (cm)", self._line("1.8", 84, "ball_cm_var"))
        self._row(card, "每行球体数量", self._spin(6, 1, 12, 84, "legend_cols_var"))
        self._row(card, "标注原子半径", self._toggle(False, "legend_radius_var"))
        self._row(card, "图例位置", self._combo(["末尾", "开头（标题后）"], 150, "legend_pos_var"))

        # ⑦ 运行
        self._section_label("⑦  运行")
        card = self._card("开始批量处理")
        run_row = QHBoxLayout()
        run_row.setContentsMargins(18, 14, 18, 6)
        run_row.setSpacing(8)
        self.run_btn = self._pill("▶  开始处理", "green", height=40, name="run_btn")
        self.pause_btn = self._pill("⏸  暂停处理", "gray", height=40, name="pause_btn")
        self.resume_btn = self._pill("▶  继续处理", "accent", height=40, name="resume_btn")
        self.stop_btn = self._pill("■  停止", "red", height=40, name="stop_btn")
        self.pause_btn.setEnabled(False)
        self.resume_btn.setEnabled(False)
        self.stop_btn.setEnabled(False)
        run_row.addWidget(self.run_btn)
        run_row.addWidget(self.pause_btn)
        run_row.addWidget(self.resume_btn)
        run_row.addWidget(self.stop_btn)
        self.progress = QProgressBar()
        self.progress.setRange(0, 1000)
        self.progress.setValue(0)
        self.progress.setTextVisible(False)
        run_row.addWidget(self.progress, 1)
        card.layout().addLayout(run_row)
        self.run_btn.clicked.connect(self._on_run)
        self.pause_btn.clicked.connect(self._on_pause)
        self.resume_btn.clicked.connect(self._on_resume)
        self.stop_btn.clicked.connect(self._on_stop)

        log_wrap = QWidget()
        log_lay = QHBoxLayout(log_wrap)
        log_lay.setContentsMargins(18, 6, 18, 16)
        self.log = QPlainTextEdit()
        self.log.setObjectName("Log")
        self.log.setReadOnly(True)
        self.log.setFixedHeight(150)
        log_lay.addWidget(self.log)
        card.layout().addWidget(log_wrap)

        self._append_log("就绪。请输入结构根目录并点击「检索」。", "info")

    # ------------------------------------------------------------------ #
    # 文件 / 目录
    # ------------------------------------------------------------------ #
    def _browse_input(self):
        d = QFileDialog.getExistingDirectory(self, "选择结构根目录")
        if d:
            self.input_root.setText(d)
            self._scan()

    def _scan(self):
        root = self.input_root.text().strip()
        if not root or not Path(root).is_dir():
            QMessageBox.warning(self, "提示", "请输入一个有效的文件夹路径。")
            return
        self.root_path = Path(root)
        self.file_table.setSortingEnabled(False)
        self.file_table.setRowCount(0)
        if self.radio_vesta.isChecked():
            found = [p for p in sorted(self.root_path.rglob("*"))
                     if p.is_file() and p.suffix.lower() == ".vesta"]
            kind = ".vesta"
        else:
            found = [p for p in sorted(self.root_path.rglob("*"))
                     if p.is_file() and p.name.upper() == "CONTCAR"]
            kind = "CONTCAR"
        for p in found:
            rel = p.relative_to(self.root_path).as_posix()
            r = self.file_table.rowCount()
            self.file_table.insertRow(r)
            it0 = QTableWidgetItem(rel)
            it0.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            it0.setCheckState(Qt.Checked)
            it0.setData(Qt.UserRole, str(p))
            self.file_table.setItem(r, 0, it0)
            it1 = QTableWidgetItem(rel)          # 默认标题 = 相对路径，可编辑
            it1.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsEditable)
            it1.setToolTip("可编辑该结构在 PPT 页面上的标题")
            self.file_table.setItem(r, 1, it1)
        self.file_table.horizontalHeader().setSortIndicator(0, Qt.AscendingOrder)
        self.file_table.setSortingEnabled(True)
        self.file_table.sortItems(0, Qt.AscendingOrder)
        self.scan_info.setText(f"检索到 {len(found)} 个 {kind}")
        self._append_log(f"在 {root} 下检索到 {len(found)} 个 {kind} 文件。", "info")
        self._refresh_elements()

    def _on_type_changed(self, checked):
        # 切换文件类型时，若已有有效路径则自动重新检索
        if checked and self.input_root.text().strip():
            self._scan()

    def _on_header_clicked(self, col):
        # 点击表头按该列排序（手动上移/下移后也能恢复排序）
        self.file_table.setSortingEnabled(True)
        self.file_table.sortItems(
            col, self.file_table.horizontalHeader().sortIndicatorOrder())

    def _row_data(self, r):
        it0 = self.file_table.item(r, 0)
        it1 = self.file_table.item(r, 1)
        return {
            "file": it0.text() if it0 else "",
            "checked": it0.checkState() if it0 else Qt.Unchecked,
            "path": it0.data(Qt.UserRole) if it0 else None,
            "title": it1.text() if it1 else "",
        }

    def _set_row_data(self, r, d):
        it0 = QTableWidgetItem(d["file"])
        it0.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
        it0.setCheckState(d["checked"])
        it0.setData(Qt.UserRole, d["path"])
        self.file_table.setItem(r, 0, it0)
        it1 = QTableWidgetItem(d["title"])
        it1.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsEditable)
        it1.setToolTip("可编辑该结构在 PPT 页面上的标题")
        self.file_table.setItem(r, 1, it1)

    def _move_row(self, delta):
        """把选中行上移(-1) / 下移(+1) 一位；手动顺序不被排序覆盖。"""
        r = self.file_table.currentRow()
        if r < 0:
            QMessageBox.information(self, "提示", "请先选中要移动的一行。")
            return
        nr = r + delta
        if nr < 0 or nr >= self.file_table.rowCount():
            return
        self.file_table.setSortingEnabled(False)
        a, b = self._row_data(r), self._row_data(nr)
        self._set_row_data(r, b)
        self._set_row_data(nr, a)
        self.file_table.setCurrentCell(nr, 1)
        self.file_table.selectRow(nr)

    def _move_to(self, to_end):
        """把选中行直接放到最前(to_end=False) / 最后(to_end=True)。"""
        r = self.file_table.currentRow()
        if r < 0:
            QMessageBox.information(self, "提示", "请先选中要移动的一行。")
            return
        n = self.file_table.rowCount()
        if n <= 1:
            return
        rows = [self._row_data(i) for i in range(n)]
        item = rows.pop(r)
        if to_end:
            rows.append(item)
        else:
            rows.insert(0, item)
        self.file_table.setSortingEnabled(False)
        for i, d in enumerate(rows):
            self._set_row_data(i, d)
        nr = n - 1 if to_end else 0
        self.file_table.setCurrentCell(nr, 1)
        self.file_table.selectRow(nr)

    def _check_all(self, checked):
        state = Qt.Checked if checked else Qt.Unchecked
        for r in range(self.file_table.rowCount()):
            it = self.file_table.item(r, 0)
            if it is not None:
                it.setCheckState(state)

    def _collect_elements_from_checked(self):
        elems = []
        for r in range(self.file_table.rowCount()):
            it = self.file_table.item(r, 0)
            if it is None or it.checkState() != Qt.Checked:
                continue
            p = Path(it.data(Qt.UserRole))
            if p.suffix.lower() == ".vesta":
                seq = [x["element"] for x in parse_atomt(p)]
            else:
                seq = parse_contcar_elements(p)
            for el in seq:
                if el not in elems:
                    elems.append(el)
        return elems

    def _refresh_elements(self):
        self._populate_atoms(self._collect_elements_from_checked())

    def _browse_outdir(self):
        d = QFileDialog.getExistingDirectory(self, "选择输出目录")
        if d:
            self.out_dir.setText(d)

    def _browse_ppt(self):
        p, _ = QFileDialog.getSaveFileName(
            self, "选择 PPT 保存路径", self.ppt_path.text(), "PowerPoint (*.pptx)")
        if p:
            self.ppt_path.setText(p)

    def _browse_vesta(self):
        p, _ = QFileDialog.getOpenFileName(
            self, "选择 VESTA.exe", "", "VESTA 可执行文件 (VESTA.exe);;所有文件 (*)")
        if p:
            self.vesta_exe_var.setText(p)

    # ------------------------------------------------------------------ #
    # 原子样式行
    # ------------------------------------------------------------------ #
    def _add_atom_row(self, element=None, radius=None, color=None):
        el = element or self.element_combo.currentText()
        radius = radius if radius is not None else f"{ELEMENTS_RADIUS.get(el, 1.4):.2f}"
        color = color or JMOL_COLORS.get(el, "#909090")

        frame = QWidget()
        h = QHBoxLayout(frame)
        h.setContentsMargins(0, 0, 0, 0)
        h.setSpacing(6)
        e = QLineEdit(el)
        e.setFixedWidth(70)
        e.setAlignment(Qt.AlignCenter)
        r = QLineEdit(radius)
        r.setFixedWidth(80)
        r.setAlignment(Qt.AlignCenter)
        c = QLineEdit(color)
        c.setFixedWidth(96)
        c.setAlignment(Qt.AlignCenter)

        # 颜色色轮：点色块打开取色器；输入框变化时同步更新
        sw = QPushButton()
        sw.setFixedSize(28, 28)
        sw.setCursor(Qt.PointingHandCursor)
        sw.setToolTip("点击打开色轮选择颜色")

        def _cur_qcolor():
            t = c.text().strip()
            qc = QColor(t)
            if qc.isValid():
                return qc
            rgb = parse_color(t)
            if rgb:
                rr, gg, bb = (int(x) for x in rgb.split())
                return QColor(rr, gg, bb)
            return QColor("#808080")

        def refresh_swatch(*_):
            bg = _cur_qcolor().name()
            sw.setStyleSheet(
                f"QPushButton{{background:{bg}; border:1px solid #D9D9DE; border-radius:7px;}}"
                "QPushButton:hover{border:1px solid #0A84FF;}")

        def pick_color():
            chosen = QColorDialog.getColor(_cur_qcolor(), self, "选择原子颜色")
            if chosen.isValid():
                c.setText(chosen.name().upper())

        sw.clicked.connect(pick_color)
        c.textChanged.connect(refresh_swatch)
        refresh_swatch()

        d = QPushButton("✕")
        d.setObjectName("DeleteBtn")
        d.setFixedSize(28, 28)
        d.setCursor(Qt.PointingHandCursor)
        d.clicked.connect(lambda: self._remove_atom_row(frame))
        h.addWidget(e)
        h.addWidget(r)
        h.addWidget(c)
        h.addWidget(sw)
        h.addWidget(d)
        h.addStretch(1)
        self.atom_rows.append({"frame": frame, "element": e, "radius": r, "color": c})
        self.atom_box_lay.addWidget(frame)

    def _remove_atom_row(self, frame):
        self.atom_rows = [row for row in self.atom_rows if row["frame"] is not frame]
        frame.deleteLater()

    def _populate_atoms(self, elements):
        for row in list(self.atom_rows):
            row["frame"].deleteLater()
        self.atom_rows.clear()
        for el in elements:
            self._add_atom_row(el)

    # ------------------------------------------------------------------ #
    # 日志 / 进度
    # ------------------------------------------------------------------ #
    def _append_log(self, msg, level="info"):
        colors = {"info": "#FFFFFF", "ok": "#30D158", "err": "#FF453A", "warn": "#FFD60A"}
        color = colors.get(level, "#FFFFFF")
        self.log.appendHtml(f'<span style="color:{color}">{_html.escape(msg)}</span>')

    def _on_log(self, level, msg):
        self._append_log(msg, level)

    def _on_progress(self, frac):
        self.progress.setValue(int(max(0.0, min(1.0, frac)) * 1000))

    # ------------------------------------------------------------------ #
    # 运行 / 停止
    # ------------------------------------------------------------------ #
    def _on_run(self):
        if self.running:
            return
        if self.file_table.rowCount() == 0:
            QMessageBox.warning(self, "提示", "请先输入结构根目录并点击「检索」。")
            return
        if not any(self.file_table.item(r, 0) is not None
                   and self.file_table.item(r, 0).checkState() == Qt.Checked
                   for r in range(self.file_table.rowCount())):
            QMessageBox.warning(self, "提示", "请至少勾选一个 CONTCAR 文件。")
            return
        if not _HAS_PPTX:
            QMessageBox.critical(
                self, "缺少 python-pptx",
                "生成 PPT 需要 python-pptx。\n\n请在命令行执行：\n"
                "    D:\\miniconda3\\envs\\chem_env\\python.exe -m pip install python-pptx\n\n"
                "然后重新打开本程序。")
            return
        vesta_exe = self.vesta_exe_var.text().strip() or DEFAULT_VESTA_EXE
        if not Path(vesta_exe).exists():
            QMessageBox.critical(self, "找不到 VESTA",
                                 f"VESTA.exe 不存在：\n{vesta_exe}\n\n请在设置里修改路径。")
            return

        settings = self._collect_settings()
        settings["vesta_exe"] = vesta_exe

        self.running = True
        self.stop_event.clear()
        self.pause_event.clear()
        self.run_btn.setEnabled(False)
        self.pause_btn.setEnabled(True)
        self.resume_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.progress.setValue(0)
        self.log.clear()
        self._append_log("开始处理…", "info")

        self.thread = QThread(self)
        self.worker = Worker(settings, self.stop_event, self.pause_event)
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.log.connect(self._on_log)
        self.worker.progress.connect(self._on_progress)
        self.worker.done.connect(self._on_done)
        self.worker.done.connect(self.thread.quit)
        self.worker.done.connect(self.worker.deleteLater)
        self.thread.finished.connect(self.thread.deleteLater)
        self.thread.start()

    def _on_pause(self):
        self.pause_event.set()
        self.pause_btn.setEnabled(False)
        self.resume_btn.setEnabled(True)
        self._append_log("已暂停，点「继续处理」恢复。", "warn")

    def _on_resume(self):
        self.pause_event.clear()
        self.pause_btn.setEnabled(True)
        self.resume_btn.setEnabled(False)
        self._append_log("继续处理…", "info")

    def _on_stop(self):
        self.stop_event.set()
        self.pause_event.clear()
        self._append_log("已请求停止，将在当前步骤结束后终止…", "warn")

    def _on_done(self, ok, msg):
        self.running = False
        self.pause_event.clear()
        self.run_btn.setEnabled(True)
        self.pause_btn.setEnabled(False)
        self.resume_btn.setEnabled(False)
        self.stop_btn.setEnabled(False)
        self._append_log(msg, "ok" if ok else "err")

    # ------------------------------------------------------------------ #
    # 收集设置
    # ------------------------------------------------------------------ #
    def _collect_settings(self):
        def num(widget, default):
            try:
                return float(widget.text())
            except Exception:
                return default

        boundary = []
        for e in self.boundary_vars:
            t = e.text().strip()
            if t:
                try:
                    boundary.append(float(t))
                except ValueError:
                    self._append_log(f"边界值 {t!r} 无效，将忽略边界设置。", "warn")
                    boundary = []
                    break
        boundary = boundary if len(boundary) == 6 else None

        atom_params = []
        for row in self.atom_rows:
            el = row["element"].text().strip()
            if not el:
                continue
            color = parse_color(row["color"].text())
            radius = row["radius"].text().strip() or None
            atom_params.append({"atom_name": el, "radius": radius, "color_RGB": color or ""})

        files = []
        for r in range(self.file_table.rowCount()):
            it = self.file_table.item(r, 0)
            if it is None or it.checkState() != Qt.Checked:
                continue
            abspath = Path(it.data(Qt.UserRole))
            rel = None
            if self.root_path is not None:
                try:
                    rel = abspath.relative_to(self.root_path)
                except ValueError:
                    rel = None
            if rel is None:
                rel = Path(abspath.name)
            rel_sub = rel.parent.as_posix()
            if rel_sub == ".":
                rel_sub = ""
            title_item = self.file_table.item(r, 1)
            title = (title_item.text().strip() if title_item is not None else "") or rel.as_posix()
            files.append({
                "path": str(abspath),
                "rel": rel_sub,
                "stem": rel.stem or rel.name or abspath.stem,
                "label": title,
            })

        return {
            "files": files,
            "out_dir": self.out_dir.text().strip(),
            "ppt_path": self.ppt_path.text().strip(),
            "img_fmt": self.img_fmt_var.currentText(),
            "scale": self.scale_var.value(),
            "timeout": self.timeout_var.value(),
            "hide_window": self.hide_window_var.isChecked(),
            "keep_vesta": self.keep_vesta_var.isChecked(),
            "vesta_exe": self.vesta_exe_var.text().strip(),
            "do_top": self.do_top_var.isChecked(),
            "do_side": self.do_side_var.isChecked(),
            "side_dir": self.side_dir_var.currentText(),
            "rot_x": self.rot_x_var.value(),
            "rot_y": self.rot_y_var.value(),
            "rot_z": self.rot_z_var.value(),
            "model_frac": float(self.model_frac_var.value()),
            "model_fit": {"裁剪留白（推荐）": "crop", "VESTA 缩放": "zoom",
                          "裁剪 + VESTA 缩放": "both"}.get(self.model_fit_var.currentText(), "crop"),
            "slide_layout": "one" if self.slide_layout_var.currentText().startswith("每页一张") else "two",
            "img_frac": float(self.img_frac_var.value()),
            "comps": "ON" if self.comps_var.isChecked() else "OFF",
            "ucolp": "ON" if self.ucolp_var.isChecked() else "OFF",
            "sbond": "ON" if self.sbond_var.isChecked() else "OFF",
            "sbond_padding": num(self.sbond_padding_var, 0.5),
            "scale_frac": num(self.scale_frac_var, 1.0),
            "x_move": num(self.x_move_var, 0.0),
            "y_move": num(self.y_move_var, 0.0),
            "boundary": boundary,
            "atom_params": atom_params if self.apply_atom_style_var.isChecked() else [],
            "legend": self.legend_var.isChecked(),
            "ball_cm": num(self.ball_cm_var, 1.8),
            "legend_cols": int(self.legend_cols_var.value()),
            "legend_radius": self.legend_radius_var.isChecked(),
            "legend_pos": self.legend_pos_var.currentText(),
        }


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------
def _set_app_id():
    """设置 AppUserModelID，使 Windows 任务栏使用本程序图标而非 python 图标。"""
    try:
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
            "moyulyy.plotCONTCAR.1")
    except Exception:
        pass


def main():
    _set_app_id()
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setFont(QFont("Segoe UI", 10))
    app.setStyleSheet(QSS)
    if ICON_PATH.exists():
        app.setWindowIcon(QIcon(str(ICON_PATH)))

    w = MainWindow()
    if ICON_PATH.exists():
        w.setWindowIcon(QIcon(str(ICON_PATH)))
    screen = app.primaryScreen().availableGeometry()
    w.move(screen.center() - w.rect().center())
    w.show()
    if os.environ.get("CONTCAR_GUI_TEST"):   # 调试用：自动退出
        QTimer.singleShot(800, app.quit)
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
